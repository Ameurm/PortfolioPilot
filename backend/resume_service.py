from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from pypdf import PdfReader


# =========================================================
# Configuration
# =========================================================

BASE_DIR = Path(__file__).resolve().parent

RESUME_DIR = BASE_DIR / "resume_data"

RESUME_SESSIONS_DIR = (
    RESUME_DIR / "sessions"
)

EMBEDDING_MODEL = (
    "sentence-transformers/all-MiniLM-L6-v2"
)

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150


# =========================================================
# Session Paths
# =========================================================

def get_session_directory(
    session_id: str,
) -> Path:

    safe_session_id = session_id.strip()

    if not safe_session_id:

        raise ValueError(
            "Session ID is required."
        )

    session_dir = (
        RESUME_SESSIONS_DIR
        / safe_session_id
    )

    return session_dir


def get_session_upload_directory(
    session_id: str,
) -> Path:

    return (
        get_session_directory(
            session_id
        )
        / "uploads"
    )


def get_session_vectorstore_directory(
    session_id: str,
) -> Path:

    return (
        get_session_directory(
            session_id
        )
        / "vectorstore"
    )


# =========================================================
# Embeddings
# =========================================================

print(
    "Loading resume embedding model..."
)

embeddings = HuggingFaceEmbeddings(
    model_name=EMBEDDING_MODEL
)

print(
    "Resume embedding model loaded."
)


# =========================================================
# PDF Extraction
# =========================================================

def extract_pdf_text(
    file_path: Path,
) -> str:

    reader = PdfReader(
        str(file_path)
    )

    pages = []

    for page in reader.pages:

        text = page.extract_text()

        if text:

            pages.append(
                text
            )

    return "\n\n".join(
        pages
    )


# =========================================================
# DOCX Extraction
# =========================================================

def extract_docx_text(
    file_path: Path,
) -> str:

    document = DocxDocument(
        str(file_path)
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            paragraphs.append(
                text
            )

    return "\n".join(
        paragraphs
    )


# =========================================================
# Resume Text Extraction
# =========================================================

def extract_resume_text(
    file_path: Path,
) -> str:

    suffix = (
        file_path.suffix.lower()
    )

    if suffix == ".pdf":

        return extract_pdf_text(
            file_path
        )

    if suffix == ".docx":

        return extract_docx_text(
            file_path
        )

    raise ValueError(
        "Unsupported resume format. "
        "Only PDF and DOCX are supported."
    )


# =========================================================
# Text Chunking
# =========================================================

def chunk_text(
    text: str,
) -> list[str]:

    normalized_text = (
        " ".join(
            text.split()
        )
    )

    chunks = []

    start = 0

    while start < len(
        normalized_text
    ):

        end = min(
            start + CHUNK_SIZE,
            len(normalized_text),
        )

        chunk = (
            normalized_text[
                start:end
            ].strip()
        )

        if chunk:

            chunks.append(
                chunk
            )

        if end >= len(
            normalized_text
        ):

            break

        start = max(
            end - CHUNK_OVERLAP,
            start + 1,
        )

    return chunks


# =========================================================
# Create Resume FAISS Index
# =========================================================

def create_resume_index(
    resume_name: str,
    text: str,
    session_id: str,
) -> dict[str, Any]:

    chunks = chunk_text(
        text
    )

    if not chunks:

        raise ValueError(
            "No readable text was found "
            "in the resume."
        )

    vectorstore_path = (
        get_session_vectorstore_directory(
            session_id
        )
    )

    vectorstore_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    documents = []

    for index, chunk in enumerate(
        chunks
    ):

        documents.append(
            Document(
                page_content=chunk,
                metadata={
                    "source": resume_name,
                    "document_type": "resume",
                    "category": "professional_experience",
                    "access_level": "public",
                    "session_id": session_id,
                    "chunk_id": (
                        f"resume-{index + 1}"
                    ),
                },
            )
        )

    vectorstore = (
        FAISS.from_documents(
            documents,
            embeddings,
        )
    )

    vectorstore.save_local(
        str(
            vectorstore_path
        )
    )

    return {
        "resume_name": resume_name,
        "chunks": len(documents),
        "vector_store": "FAISS",
        "embedding_model": (
            EMBEDDING_MODEL
        ),
        "session_id": session_id,
    }


# =========================================================
# Process Resume
# =========================================================

def process_resume(
    file_path: Path,
    resume_name: str,
    session_id: str,
) -> dict[str, Any]:

    text = extract_resume_text(
        file_path
    )

    if not text.strip():

        raise ValueError(
            "The resume does not contain "
            "extractable text."
        )

    result = create_resume_index(
        resume_name=resume_name,
        text=text,
        session_id=session_id,
    )

    result[
        "characters_extracted"
    ] = len(text)

    return result


# =========================================================
# Resume Retriever
# =========================================================

def get_resume_retriever(
    session_id: str,
):

    vectorstore_path = (
        get_session_vectorstore_directory(
            session_id
        )
    )

    index_file = (
        vectorstore_path
        / "index.faiss"
    )

    if not index_file.exists():

        return None

    vectorstore = FAISS.load_local(
        str(
            vectorstore_path
        ),
        embeddings,
        allow_dangerous_deserialization=True,
    )

    return vectorstore.as_retriever(
        search_kwargs={
            "k": 10
        }
    )
